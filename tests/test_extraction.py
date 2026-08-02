from __future__ import annotations

import shutil

import pytest
from openpyxl import Workbook
from PIL import Image, ImageDraw, ImageFont

from taxsentry.extraction import extract


def test_xlsx_extraction_uses_existing_financial_parser(tmp_path):
    path = tmp_path / "report.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Chỉ tiêu", "Tháng 4", "Tháng 5"])
    sheet.append(["Doanh thu", 100_000_000, 120_000_000])
    sheet.append(["Lợi nhuận ròng", 20_000_000, 18_000_000])
    workbook.save(path)
    result = extract(path, ["vie", "eng"])
    assert result.source == "xlsx"
    assert result.confidence == 1.0
    assert result.content["data"]["canonical_metrics"]


def test_xlsx_extraction_prefers_current_period_and_net_revenue(tmp_path):
    path = tmp_path / "quarter.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["CHỈ TIÊU", "Mã số", "Kỳ Này", "Kỳ Trước"])
    sheet.append(["Doanh thu bán hàng và cung cấp dịch vụ", 1, 600, 500])
    sheet.append(["Doanh thu thuần", 10, 580, 480])
    sheet.append(["Doanh thu hoạt động tài chính", 21, 7, 9])
    sheet.append(["Giá vốn hàng bán", 11, 300, 250])
    workbook.save(path)

    data = extract(path, ["vie", "eng"]).content["data"]
    assert data["canonical_metrics"]["revenue"]["value"] == 580
    assert data["sheets"][0]["summary"]["revenue"]["value"] == 580
    assert data["income_statement"]["T4_Actual"] == {}
    assert data["income_statement"]["T5_Actual"] == {}


def test_xlsx_resolver_keeps_statement_period_and_rejects_decoys(tmp_path):
    path = tmp_path / "anthropic-like.xlsx"
    workbook = Workbook()
    statement = workbook.active
    statement.title = "03_KQKD"
    statement.append(["Chỉ tiêu", "FY2025A", "FY2026E"])
    statement.append(["Đơn vị: USD billion", None, None])
    statement.append(["Doanh thu", 40, 43.5])
    statement.append(["Giá vốn", -20, -25.407])
    statement.append(["Lợi nhuận gộp", 20, 18.093])
    statement.append(["Lợi nhuận trước thuế", -8, -10.423])
    statement.append(["Lợi nhuận sau thuế", -8, -10.423])
    statement.append(["Net debt", -28, -28])
    statement.append(["COGS %", -0.584, -0.584])
    scenario = workbook.create_sheet("07_Kich_ban")
    scenario.append(["Metric", "Base", "Bull"])
    scenario.append(["Revenue", 43.5, 52])
    valuation = workbook.create_sheet("Valuation")
    valuation.append(["Revenue multiple", 10])
    workbook.save(path)

    metrics = extract(path, ["vie", "eng"]).content["data"]["canonical_metrics"]
    assert metrics["revenue"]["value"] == 43.5
    assert metrics["revenue"]["source_type"] == "income_statement"
    assert metrics["revenue"]["currency"] == "USD"
    assert metrics["revenue"]["scale_multiplier"] == "1000000000"
    assert metrics["net_income"]["value"] == -10.423
    assert metrics["ebt"]["value"] == -10.423
    assert metrics["cogs"]["value"] == -25.407


@pytest.mark.skipif(not shutil.which("tesseract"), reason="Tesseract is a system prerequisite")
def test_real_image_ocr(tmp_path):
    path = tmp_path / "scan.png"
    image = Image.new("RGB", (900, 180), "white")
    ImageDraw.Draw(image).text((30, 50), "REVENUE 120000 PROFIT 18000", fill="black", font=ImageFont.load_default(size=38))
    image.save(path)
    result = extract(path, ["eng"])
    assert result.source == "ocr"
    assert "120000" in result.content
