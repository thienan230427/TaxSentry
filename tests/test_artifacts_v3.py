import pdfplumber
from docx import Document
from openpyxl import load_workbook

from taxsentry.artifacts import ArtifactSpec, render_artifact
from taxsentry.reporting import normalize_report_v3


def _v3_report():
    return normalize_report_v3({
        "schema_version": 2,
        "profile": "tax_risk_memo",
        "decision_question": "Cash có khớp không?",
        "period": {"label": "FY2026E", "start": "", "end": ""},
        "executive_summary": "Đối chiếu cần review.",
        "metrics": [{"id": "revenue", "label": "Revenue", "current": 43_500_000_000, "previous": None, "budget": None, "benchmark": None, "unit": "USD", "source_ids": ["s"], "assessment": ""}],
        "findings": [], "scenario_model": {"model_type": "none", "drivers": [], "primary_output": "", "scenarios": []}, "tax_risks": [], "recommendations": [],
        "sources": [{"id": "s", "kind": "file", "title": "statement", "locator": "sheet=PL;cell=C5", "fetched_at": "", "effective_from": "", "verified_current": True}], "missing_data": [], "assumptions": [], "overall_confidence": 0.8,
    })


def _complete_v3_report():
    report = _v3_report()
    base = report["metrics"][0]["current"]
    for index in range(1, 8):
        current = {**base, "raw_value": str(index), "normalized_value": str(index * 1_000_000_000)}
        report["metrics"].append(
            {
                "id": f"metric_{index}",
                "label": f"Metric {index}",
                "current": current,
                "previous": None,
                "budget": None,
                "benchmark": None,
                "assessment": "",
            }
        )
    report["reconciliation_checks"] = [
        {"id": "cash_flow.ending_cash_equals_balance_sheet_cash", "status": "PASS", "delta": "0", "evidence_ids": ["s"], "where_to_fix": "Cash flow"},
        {"id": "cash_flow.beginning_plus_movement_equals_ending", "status": "PASS", "delta": "0", "evidence_ids": ["s"], "where_to_fix": "Cash flow"},
        {"id": "pnl.revenue_plus_cogs_equals_gross_profit", "status": "PASS", "delta": "0", "evidence_ids": ["s"], "where_to_fix": "P&L"},
    ]
    report["scenario_model"]["scenarios"] = [
        {"name": name, "assumptions": "Documented driver and sensitivity range."}
        for name in ("base", "bull", "bear")
    ]
    report["tax_risks"] = [{"title": "Tax documentation", "severity": "medium", "regulation": ""}] * 3
    report["recommendations"] = [{"priority": "high", "owner": "CFO", "deadline_days": 30, "action": "Close the evidence gap."}] * 3
    report["missing_data"] = []
    report["data_quality"] = {"coverage": 1.0, "conflict_count": 0, "missing_material_fields": [], "confidence_caps": [], "review_reasons": []}
    return report


def test_v3_artifacts_use_audit_contract(tmp_path):
    report = _v3_report()
    report["reconciliation_checks"] = [{"id": "cash_flow.ending_cash_equals_balance_sheet_cash", "status": "FAIL", "actual": "20.859", "expected": "8.918", "delta": "11.941", "tolerance": "0.000001", "evidence_ids": ["s"], "where_to_fix": "Cash flow"}]
    spec = ArtifactSpec.from_report(report, currency="USD")
    xlsx = render_artifact("xlsx", spec, tmp_path)
    assert load_workbook(xlsx, data_only=False).sheetnames == ["Summary", "Normalized_Facts", "Checks", "Scenarios", "Sources_Audit"]
    docx = render_artifact("docx", spec, tmp_path)
    xml = Document(docx)._element.xml
    assert 'TOC \\o "1-3"' not in xml
    assert "Reconciliation" in xml
    pdf = render_artifact("pdf", spec, tmp_path)
    assert pdf.stat().st_size > 0


def test_v3_complete_pdf_uses_audit_page_budget(tmp_path):
    report = _complete_v3_report()
    spec = ArtifactSpec.from_report(report, currency="USD")
    pdf = render_artifact("pdf", spec, tmp_path)
    with pdfplumber.open(pdf) as document:
        assert 8 <= len(document.pages) <= 12


def test_pdf_and_docx_use_distinct_content_maps(tmp_path):
    report = _v3_report()
    spec = ArtifactSpec.from_report(report, currency="USD")
    docx = render_artifact("docx", spec, tmp_path)
    headings = {paragraph.text for paragraph in Document(docx).paragraphs}
    assert "Data inventory" in headings
    assert "Full source lineage" in headings

    pdf = render_artifact("pdf", spec, tmp_path)
    with pdfplumber.open(pdf) as document:
        pdf_text = "\n".join(page.extract_text() or "" for page in document.pages)
    assert "Executive verdict" in pdf_text
    assert "Data inventory" not in pdf_text
