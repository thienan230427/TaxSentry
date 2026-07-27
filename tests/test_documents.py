from __future__ import annotations

import asyncio
import hashlib
import json
import zipfile

import pytest
from docx import Document
from openpyxl import Workbook
from pptx import Presentation

from taxsentry.documents import (
    DocumentService,
    DocumentValidationError,
    ParsedDocument,
    _split_complete,
    document_service_from_settings,
    validate_document,
)
from taxsentry.extraction import extract


def test_document_service_uses_distributed_coordinator_when_enabled(monkeypatch):
    from taxsentry.data_plane import DistributedDocumentService

    marker = object()
    monkeypatch.setattr(
        DistributedDocumentService,
        "from_settings",
        classmethod(lambda cls, settings: marker),
    )

    assert document_service_from_settings({"data_plane": {"enabled": True}}) is marker
    assert isinstance(document_service_from_settings({}), DocumentService)


def test_excel_manifest_covers_every_sheet_row_and_locator(tmp_path):
    path = tmp_path / "case.xlsx"
    workbook = Workbook()
    visible = workbook.active
    visible.title = "Kỳ Này"
    visible.append(["Chỉ tiêu", "Kỳ Này", "Kỳ Trước"])
    visible.append(["Doanh thu thuần", 580, 480])
    hidden = workbook.create_sheet("Tổng hợp thuế")
    hidden.sheet_state = "veryHidden"
    hidden.append(["Chỉ tiêu", "Số tiền"])
    hidden.append(["BHXH", 12])
    formula = workbook.create_sheet("Formula")
    formula["A1"] = "=SUM('Kỳ Này'!B2:C2)"
    workbook.save(path)

    service = DocumentService(tmp_path / "documents")
    manifest = asyncio.run(
        service.ingest(path=path, company_id="acme", case_id="case-1")
    )

    assert manifest.coverage.complete
    assert manifest.metadata["sheets"][1]["state"] == "veryHidden"
    assert {item["name"] for item in manifest.metadata["sheets"]} == {
        "Kỳ Này",
        "Tổng hợp thuế",
        "Formula",
    }
    assert (
        asyncio.run(service.coverage(manifest.id, company_id="acme")).processed
        == manifest.coverage.total
    )
    hit = asyncio.run(
        service.query(
            document_id=manifest.id,
            company_id="acme",
            query="Doanh thu thuần",
            limit=1,
        )
    )[0]
    assert hit.locator.startswith("sheet=Kỳ Này!")
    assert "B2=580" in hit.text
    assert any(
        unit.structured.get("sheet") == "Formula"
        and unit.structured["rows"][0]["formula_dependencies"]
        for unit in service.iter_units(manifest.id, company_id="acme")
        if unit.kind == "sheet_rows"
    )
    compatibility = extract(path, ["vie", "eng"]).content
    assert compatibility["data"]["canonical_metrics"]["revenue"]["value"] == 580
    assert compatibility["data"]["sheets"][1]["type"] == "tax_summary"


def test_word_and_powerpoint_are_structural_units(tmp_path):
    word_path = tmp_path / "memo.docx"
    document = Document()
    document.sections[0].header.paragraphs[0].text = "ACME confidential"
    document.add_heading("Kết quả", 1)
    document.add_paragraph("Doanh thu tăng.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Doanh thu"
    table.rows[0].cells[1].text = "580"
    document.save(word_path)

    ppt_path = tmp_path / "deck.pptx"
    presentation = Presentation()
    for index in range(12):
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        slide.shapes.title.text = f"Slide {index + 1}"
    presentation.save(ppt_path)

    service = DocumentService(tmp_path / "documents")
    word = service.ingest_sync(path=word_path, company_id="acme")
    deck = service.ingest_sync(path=ppt_path, company_id="acme")

    word_units = list(service.iter_units(word.id, company_id="acme"))
    slide_units = [
        unit
        for unit in service.iter_units(deck.id, company_id="acme")
        if unit.kind == "slide"
    ]
    assert {"paragraph", "table", "header"} <= {unit.kind for unit in word_units}
    assert [unit.locator for unit in slide_units] == [
        f"slide={index}" for index in range(1, 13)
    ]


def test_prompt_partition_drops_no_characters():
    source = "0123456789" * 1_200
    parts = list(_split_complete(source, 1_000))
    rebuilt = "".join(part.split("\n", 1)[1] for part in parts)
    assert rebuilt == source


def test_analysis_payload_is_not_silently_dropped_above_legacy_limit(tmp_path):
    source = tmp_path / "large.pdf"
    source.write_bytes(b"%PDF-1.7\n")
    payload = "x" * (20 * 1024 * 1024 + 1)
    service = DocumentService(tmp_path / "documents")
    manifest = service._persist(
        source,
        ".pdf",
        hashlib.sha256(source.read_bytes()).hexdigest(),
        "d" * 32,
        "acme",
        "case-1",
        ParsedDocument("fixture", (), {"payload": payload}, 1.0),
    )

    assert len(
        service.analysis_content(manifest.id, company_id="acme")["payload"]
    ) == len(payload)


def test_rejects_mime_spoof_and_zip_traversal(tmp_path):
    fake = tmp_path / "fake.pdf"
    fake.write_bytes(b"not a pdf")
    with pytest.raises(DocumentValidationError, match="signature"):
        validate_document(fake)

    archive_path = tmp_path / "bad.docx"
    with zipfile.ZipFile(archive_path, "w") as archive:
        archive.writestr("../escape.xml", json.dumps({"bad": True}))
    with pytest.raises(DocumentValidationError, match="traversal"):
        validate_document(archive_path)

    bomb_path = tmp_path / "bomb.docx"
    with zipfile.ZipFile(bomb_path, "w") as archive:
        archive.writestr(
            "filler.bin",
            b"x" * (2 * 1024 * 1024),
            compress_type=zipfile.ZIP_STORED,
        )
        archive.writestr(
            "bomb.bin",
            b"\x00" * (2 * 1024 * 1024),
            compress_type=zipfile.ZIP_DEFLATED,
        )
    with pytest.raises(DocumentValidationError, match="member compression ratio"):
        validate_document(bomb_path)
