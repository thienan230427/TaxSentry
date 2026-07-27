from __future__ import annotations

import argparse
import asyncio
import json
import math
import tempfile
import time
import zipfile
from collections import Counter
from dataclasses import asdict, dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Sequence

from docx import Document
from openpyxl import Workbook
from PIL import Image, ImageDraw
from pptx import Presentation
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

from taxsentry.documents import EXCEL_ROWS_PER_UNIT, DocumentService


@dataclass(frozen=True)
class BenchmarkProfile:
    xlsx_sheets: int
    xlsx_rows_per_sheet: int
    pdf_pages: int
    pdf_scan_every: int
    pptx_slides: int
    docx_paragraphs: int
    docx_table_rows: int
    target_case_mb: int = 0


PROFILES = {
    "quick": BenchmarkProfile(3, 20, 4, 0, 5, 10, 4),
    "acceptance": BenchmarkProfile(
        100,
        10_000,
        1_000,
        10,
        500,
        2_000,
        200,
        497,
    ),
}


@dataclass(frozen=True)
class GeneratedFixture:
    name: str
    path: Path
    generation_seconds: float
    expected_minimum_units: dict[str, int]
    metadata: dict[str, Any]


def run_benchmark(profile: str = "quick", root: Path | None = None) -> dict[str, Any]:
    if profile not in PROFILES:
        raise ValueError(f"Unknown profile: {profile}")
    if root is None:
        with tempfile.TemporaryDirectory(prefix=f"taxsentry-{profile}-benchmark-") as folder:
            return _run(profile, Path(folder))
    root.mkdir(parents=True, exist_ok=True)
    return _run(profile, root)


def _run(profile: str, root: Path) -> dict[str, Any]:
    spec = PROFILES[profile]
    fixture_root = root / "fixtures"
    fixture_root.mkdir(parents=True, exist_ok=True)
    fixtures = [
        _generate_xlsx(fixture_root / "large.xlsx", spec),
        _generate_pdf(fixture_root / "large.pdf", spec),
        _generate_pptx(fixture_root / "large.pptx", spec),
        _generate_docx(fixture_root / "large.docx", spec),
    ]
    service = DocumentService(root / "documents")
    document_results: list[dict[str, Any]] = []
    for fixture in fixtures:
        started = time.perf_counter()
        manifest = service.ingest_sync(
            path=fixture.path,
            company_id="benchmark-company",
            case_id=f"benchmark-{profile}",
            languages=["eng"],
        )
        elapsed = time.perf_counter() - started
        units = list(
            service.iter_units(
                manifest.id,
                company_id="benchmark-company",
            )
        )
        observed = Counter(unit.kind for unit in units)
        missing = {
            kind: expected - observed.get(kind, 0)
            for kind, expected in fixture.expected_minimum_units.items()
            if observed.get(kind, 0) < expected
        }
        document_results.append(
            {
                "name": fixture.name,
                "path": str(fixture.path),
                "size_bytes": fixture.path.stat().st_size,
                "generation_seconds": round(fixture.generation_seconds, 6),
                "ingest_seconds": round(elapsed, 6),
                "coverage": asdict(manifest.coverage),
                "observed_units": dict(sorted(observed.items())),
                "expected_minimum_units": fixture.expected_minimum_units,
                "missing_units": missing,
                "no_silent_unit_loss": not missing
                and len(units) == manifest.coverage.total,
                "metadata": fixture.metadata,
                "document_id": manifest.id,
            }
        )

    case_started = time.perf_counter()
    case = asyncio.run(
        service.ingest_case(
            paths=[fixture.path for fixture in fixtures],
            company_id="benchmark-company",
            case_id=f"benchmark-{profile}",
            languages=["eng"],
        )
    )
    case_elapsed = time.perf_counter() - case_started
    total = sum(item["coverage"]["total"] for item in document_results)
    processed = sum(item["coverage"]["processed"] for item in document_results)
    failed = sum(item["coverage"]["failed"] for item in document_results)
    skipped = sum(item["coverage"]["skipped"] for item in document_results)
    passed = (
        case.coverage_complete
        and total == processed
        and not failed
        and not skipped
        and all(item["no_silent_unit_loss"] for item in document_results)
    )
    return {
        "schema_version": 1,
        "profile": profile,
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "settings": asdict(spec),
        "documents": document_results,
        "case": {
            "id": case.id,
            "document_ids": list(case.document_ids),
            "size_bytes": case.size_bytes,
            "ingest_seconds": round(case_elapsed, 6),
            "coverage": {
                "total": total,
                "processed": processed,
                "failed": failed,
                "skipped": skipped,
                "complete": case.coverage_complete and total == processed,
            },
        },
        "passed": passed,
    }


def _generate_xlsx(path: Path, spec: BenchmarkProfile) -> GeneratedFixture:
    started = time.perf_counter()
    workbook = Workbook(write_only=True)
    for sheet_number in range(1, spec.xlsx_sheets + 1):
        sheet = workbook.create_sheet(f"Sheet{sheet_number:03d}")
        if sheet_number > 1 and sheet_number % 10 == 0:
            sheet.sheet_state = "hidden"
        sheet.append(("row_id", "amount", "tax"))
        for row_number in range(1, spec.xlsx_rows_per_sheet + 1):
            tax: int | str = (
                "=B2*0.1" if row_number == 1 else row_number % 17
            )
            sheet.append((row_number, row_number * 10, tax))
    workbook.save(path)
    rows = spec.xlsx_rows_per_sheet + 1
    _add_xlsx_dimensions(path, rows)
    if spec.target_case_mb:
        _pad_ooxml(path, spec.target_case_mb * 1024 * 1024)
    expected_chunks = spec.xlsx_sheets * math.ceil(
        rows / EXCEL_ROWS_PER_UNIT
    )
    return GeneratedFixture(
        "xlsx",
        path,
        time.perf_counter() - started,
        {"workbook": 1, "sheet_rows": expected_chunks},
        {
            "sheets": spec.xlsx_sheets,
            "data_rows": spec.xlsx_sheets * spec.xlsx_rows_per_sheet,
            "total_rows_including_headers": spec.xlsx_sheets * rows,
        },
    )


def _add_xlsx_dimensions(path: Path, rows: int) -> None:
    """openpyxl write-only sheets omit dimensions; add them without loading 1M rows."""

    temporary = path.with_name(f".{path.stem}-dimensions.xlsx")
    with zipfile.ZipFile(path) as source, zipfile.ZipFile(temporary, "w") as target:
        for info in source.infolist():
            payload = source.read(info)
            if info.filename.startswith("xl/worksheets/sheet") and info.filename.endswith(
                ".xml"
            ):
                root_end = payload.find(b">")
                if root_end < 0:
                    raise ValueError(f"Invalid worksheet XML: {info.filename}")
                dimension = f'<dimension ref="A1:C{rows}"/>'.encode()
                payload = payload[: root_end + 1] + dimension + payload[root_end + 1 :]
            target.writestr(info, payload)
    temporary.replace(path)


def _pad_ooxml(path: Path, target_bytes: int) -> None:
    """Exercise the near-500 MB transport path without adding fake worksheet rows."""

    padding = target_bytes - path.stat().st_size - 512
    if padding <= 0:
        return
    chunk = bytes(1024 * 1024)
    with zipfile.ZipFile(
        path,
        "a",
        compression=zipfile.ZIP_STORED,
        allowZip64=True,
    ) as archive:
        with archive.open(
            "xl/media/taxsentry-benchmark-padding.bin",
            "w",
            force_zip64=True,
        ) as output:
            while padding:
                size = min(padding, len(chunk))
                output.write(chunk[:size])
                padding -= size


def _generate_pdf(path: Path, spec: BenchmarkProfile) -> GeneratedFixture:
    started = time.perf_counter()
    scan_path = path.with_name("scan-page.png")
    image = Image.new("RGB", (800, 1_100), "white")
    draw = ImageDraw.Draw(image)
    scan_text = (
        "Scanned invoice evidence for TaxSentry benchmark. "
        "Revenue expense tax accounting source verification. "
    )
    for line in range(12):
        draw.text((40, 40 + line * 70), f"{line + 1}. {scan_text}", fill="black")
    image.save(scan_path)

    pdf = canvas.Canvas(str(path), pagesize=letter, pageCompression=1)
    width, height = letter
    scan_pages = 0
    for page_number in range(1, spec.pdf_pages + 1):
        if spec.pdf_scan_every and page_number % spec.pdf_scan_every == 0:
            pdf.drawImage(
                str(scan_path),
                0,
                0,
                width=width,
                height=height,
                preserveAspectRatio=True,
            )
            scan_pages += 1
        else:
            text = pdf.beginText(36, height - 48)
            for line in range(5):
                text.textLine(
                    f"Page {page_number}, line {line + 1}. "
                    "TaxSentry benchmark revenue expense tax accounting evidence "
                    "with enough text to use deterministic PDF extraction."
                )
            pdf.drawText(text)
        pdf.showPage()
    pdf.save()
    scan_path.unlink(missing_ok=True)
    return GeneratedFixture(
        "pdf",
        path,
        time.perf_counter() - started,
        {"page": spec.pdf_pages},
        {
            "pages": spec.pdf_pages,
            "text_pages": spec.pdf_pages - scan_pages,
            "simulated_scan_pages": scan_pages,
        },
    )


def _generate_pptx(path: Path, spec: BenchmarkProfile) -> GeneratedFixture:
    started = time.perf_counter()
    presentation = Presentation()
    for slide_number in range(1, spec.pptx_slides + 1):
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        slide.shapes.title.text = f"Benchmark slide {slide_number}"
    presentation.save(path)
    return GeneratedFixture(
        "pptx",
        path,
        time.perf_counter() - started,
        {"slide": spec.pptx_slides},
        {"slides": spec.pptx_slides},
    )


def _generate_docx(path: Path, spec: BenchmarkProfile) -> GeneratedFixture:
    started = time.perf_counter()
    document = Document()
    document.sections[0].header.paragraphs[0].text = "TaxSentry benchmark header"
    document.sections[0].footer.paragraphs[0].text = "TaxSentry benchmark footer"
    document.add_heading("Benchmark report", 1)
    commented = document.add_paragraph(
        "Grounded financial analysis with source citations."
    )
    for paragraph_number in range(1, spec.docx_paragraphs):
        document.add_paragraph(
            f"Paragraph {paragraph_number}: revenue expense tax accounting evidence."
        )
    table = document.add_table(rows=1, cols=3)
    for cell, value in zip(
        table.rows[0].cells,
        ("row_id", "amount", "tax"),
        strict=True,
    ):
        cell.text = value
    for row_number in range(1, spec.docx_table_rows + 1):
        cells = table.add_row().cells
        cells[0].text = str(row_number)
        cells[1].text = str(row_number * 10)
        cells[2].text = str(row_number % 17)
    comment_supported = hasattr(document, "add_comment")
    if comment_supported:
        document.add_comment(
            commented.runs,
            text="Benchmark comment",
            author="TaxSentry",
            initials="TS",
        )
    document.save(path)
    expected = {
        "paragraph": spec.docx_paragraphs + 1,
        "table": 1,
        "header": 1,
        "footer": 1,
    }
    if comment_supported:
        expected["comment"] = 1
    return GeneratedFixture(
        "docx",
        path,
        time.perf_counter() - started,
        expected,
        {
            "paragraphs": spec.docx_paragraphs + 1,
            "table_rows": spec.docx_table_rows + 1,
            "comment_supported": comment_supported,
        },
    )


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description="Generate temporary large-document fixtures and benchmark ingestion."
    )
    parser.add_argument(
        "--profile",
        choices=tuple(PROFILES),
        default="quick",
    )
    parser.add_argument("--output", type=Path)
    args = parser.parse_args(argv)
    result = run_benchmark(args.profile)
    encoded = json.dumps(result, ensure_ascii=False, indent=2) + "\n"
    if args.output:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(encoded, encoding="utf-8")
    else:
        print(encoded, end="")
    return 0 if result["passed"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
