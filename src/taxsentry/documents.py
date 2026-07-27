from __future__ import annotations

import asyncio
import hashlib
import json
import os
import re
import shutil
import subprocess
import tempfile
import zipfile
from dataclasses import asdict, dataclass, field, replace
from datetime import date, datetime, timezone
from pathlib import Path, PurePosixPath
from typing import Any, Callable, Iterable, Iterator, Mapping
from xml.etree import ElementTree

from .config import APP_HOME
from .data_plane import LocalObjectStore, ObjectStore

MAX_DOCUMENT_BYTES = 500 * 1024 * 1024
MAX_ARCHIVE_BYTES = 4 * 1024 * 1024 * 1024
MAX_ARCHIVE_RATIO = 1_000
MAX_ARCHIVE_ENTRIES = 100_000
EXCEL_ROWS_PER_UNIT = 250
DOCUMENT_EXTRACTION_SCHEMA_VERSION = 1
PROMPT_PARTITION_CHARS = 60_000
EVIDENCE_SCHEMA: dict[str, Any] = {
    "type": "object",
    "properties": {
        "summary": {"type": "string"},
        "claims": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {
                    "statement": {"type": "string"},
                    "source_ids": {"type": "array", "items": {"type": "string"}},
                },
                "required": ["statement", "source_ids"],
                "additionalProperties": False,
            },
        },
        "conflicts": {"type": "array", "items": {"type": "string"}},
        "missing_data": {"type": "array", "items": {"type": "string"}},
    },
    "required": ["summary", "claims", "conflicts", "missing_data"],
    "additionalProperties": False,
}
SUPPORTED_SUFFIXES = {
    ".doc",
    ".docx",
    ".xls",
    ".xlsx",
    ".xlsm",
    ".ppt",
    ".pptx",
    ".pdf",
    ".png",
    ".jpg",
    ".jpeg",
}
OOXML_SUFFIXES = {".docx", ".xlsx", ".xlsm", ".pptx"}
LEGACY_SUFFIXES = {".doc", ".xls", ".ppt"}
_OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
_FORMULA_REF = re.compile(
    r"(?:(?:'([^']+)'|([A-Za-z0-9_. -]+))!)?\$?([A-Z]{1,3})\$?([1-9][0-9]*)"
)
_TOKEN = re.compile(r"[\wÀ-ỹ]+", re.UNICODE)


class DocumentValidationError(ValueError):
    pass


class DocumentNeedsReview(ValueError):
    pass


@dataclass(frozen=True, slots=True)
class EvidenceRef:
    source_id: str
    locator: str
    unit_type: str
    ordinal: int
    bounding_box: tuple[float, float, float, float] | None = None


@dataclass(slots=True)
class DocumentUnit:
    id: str
    document_id: str
    kind: str
    ordinal: int
    locator: str
    text: str = ""
    structured: dict[str, Any] = field(default_factory=dict)
    status: str = "processed"
    error: str = ""
    warnings: list[str] = field(default_factory=list)

    @property
    def evidence(self) -> EvidenceRef:
        return EvidenceRef(self.id, self.locator, self.kind, self.ordinal)


@dataclass(frozen=True, slots=True)
class CoverageReport:
    total: int
    processed: int
    failed: int
    skipped: int = 0
    failed_units: tuple[str, ...] = ()
    warnings: tuple[str, ...] = ()

    @property
    def complete(self) -> bool:
        return self.total == self.processed and not self.failed and not self.skipped


@dataclass(slots=True)
class DocumentManifest:
    id: str
    company_id: str
    case_id: str
    name: str
    suffix: str
    sha256: str
    size_bytes: int
    source: str
    raw_uri: str
    units_uri: str
    created_at: str
    coverage: CoverageReport
    metadata: dict[str, Any] = field(default_factory=dict)


@dataclass(frozen=True, slots=True)
class CaseManifest:
    id: str
    company_id: str
    document_ids: tuple[str, ...]
    size_bytes: int
    coverage_complete: bool


@dataclass(slots=True)
class ParsedDocument:
    source: str
    units: Iterable[DocumentUnit]
    compatibility_content: dict[str, Any] | str
    confidence: float
    metadata: dict[str, Any] = field(default_factory=dict)


def validate_document(path: Path, *, max_bytes: int = MAX_DOCUMENT_BYTES) -> str:
    resolved = path.expanduser().resolve()
    if not resolved.is_file():
        raise FileNotFoundError(resolved)
    suffix = resolved.suffix.casefold()
    if suffix not in SUPPORTED_SUFFIXES:
        raise DocumentValidationError(f"Unsupported document type: {suffix or '(none)'}")
    size = resolved.stat().st_size
    if not size:
        raise DocumentValidationError("Document is empty")
    if size > max_bytes:
        raise DocumentValidationError(
            f"Document exceeds the {max_bytes // (1024 * 1024)} MB limit"
        )
    with resolved.open("rb") as stream:
        magic = stream.read(1_024)
    if suffix in OOXML_SUFFIXES:
        if not magic.startswith((b"PK\x03\x04", b"PK\x05\x06", b"PK\x07\x08")):
            raise DocumentValidationError("OOXML extension does not match file signature")
        _validate_archive(resolved)
    elif suffix in LEGACY_SUFFIXES and not magic.startswith(_OLE_MAGIC):
        raise DocumentValidationError("Legacy Office extension does not match OLE signature")
    elif suffix == ".pdf" and b"%PDF-" not in magic:
        raise DocumentValidationError("PDF extension does not match file signature")
    elif suffix == ".png" and not magic.startswith(b"\x89PNG\r\n\x1a\n"):
        raise DocumentValidationError("PNG extension does not match file signature")
    elif suffix in {".jpg", ".jpeg"} and not magic.startswith(b"\xff\xd8\xff"):
        raise DocumentValidationError("JPEG extension does not match file signature")
    return suffix


def _validate_archive(path: Path) -> None:
    try:
        with zipfile.ZipFile(path) as archive:
            infos = archive.infolist()
            if len(infos) > MAX_ARCHIVE_ENTRIES:
                raise DocumentValidationError("Archive contains too many entries")
            expanded = 0
            compressed = 0
            for info in infos:
                member = PurePosixPath(info.filename.replace("\\", "/"))
                if member.is_absolute() or ".." in member.parts:
                    raise DocumentValidationError("Archive path traversal detected")
                expanded += info.file_size
                compressed += max(1, info.compress_size)
                if (
                    info.file_size
                    and info.file_size / max(1, info.compress_size)
                    > MAX_ARCHIVE_RATIO
                ):
                    raise DocumentValidationError(
                        "Suspicious archive member compression ratio"
                    )
                if expanded > MAX_ARCHIVE_BYTES:
                    raise DocumentValidationError("Archive expands beyond the safety limit")
            if expanded and expanded / max(1, compressed) > MAX_ARCHIVE_RATIO:
                raise DocumentValidationError("Suspicious archive compression ratio")
    except zipfile.BadZipFile as exc:
        raise DocumentValidationError("Corrupt OOXML archive") from exc


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _unit_id(document_id: str, kind: str, ordinal: int, locator: str) -> str:
    seed = f"{document_id}:{kind}:{ordinal}:{locator}".encode()
    return hashlib.sha256(seed).hexdigest()[:32]


def _new_unit(
    document_id: str,
    kind: str,
    ordinal: int,
    locator: str,
    *,
    text: str = "",
    structured: dict[str, Any] | None = None,
    status: str = "processed",
    error: str = "",
    warnings: list[str] | None = None,
) -> DocumentUnit:
    return DocumentUnit(
        _unit_id(document_id, kind, ordinal, locator),
        document_id,
        kind,
        ordinal,
        locator,
        text.strip(),
        structured or {},
        status,
        error,
        warnings or [],
    )


def _json_value(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return str(value)


def _text_value(value: Any) -> str:
    value = _json_value(value)
    return "" if value is None else str(value)


def parse_document(
    path: Path,
    document_id: str,
    languages: list[str],
    *,
    excel_rows_per_unit: int = EXCEL_ROWS_PER_UNIT,
) -> ParsedDocument:
    suffix = validate_document(path)
    if suffix in LEGACY_SUFFIXES:
        return _legacy_document(
            path,
            document_id,
            languages,
            excel_rows_per_unit=excel_rows_per_unit,
        )
    if suffix in {".xlsx", ".xlsm"}:
        return _excel_document(
            path,
            document_id,
            rows_per_unit=excel_rows_per_unit,
        )
    if suffix == ".docx":
        return _word_document(path, document_id, languages)
    if suffix == ".pptx":
        return _powerpoint_document(path, document_id, languages)
    if suffix == ".pdf":
        return _pdf_document(path, document_id, languages)
    return _image_document(path, document_id, languages)


class DocumentService:
    """Persistent structural document ingestion with no silent unit truncation."""

    def __init__(
        self,
        root: Path | None = None,
        object_store: ObjectStore | None = None,
        *,
        excel_rows_per_unit: int = EXCEL_ROWS_PER_UNIT,
        progress_callback: Callable[[DocumentUnit, int], None] | None = None,
    ):
        if excel_rows_per_unit < 1:
            raise ValueError("excel_rows_per_unit must be positive")
        self.root = (root or APP_HOME / "documents").expanduser().resolve()
        self.object_store = object_store or LocalObjectStore(
            self.root.parent / "objects"
        )
        self.excel_rows_per_unit = int(excel_rows_per_unit)
        self.progress_callback = progress_callback

    async def ingest(
        self,
        *,
        path: Path,
        company_id: str,
        case_id: str = "",
        languages: list[str] | None = None,
    ) -> DocumentManifest:
        return await asyncio.to_thread(
            self.ingest_sync,
            path=path,
            company_id=company_id,
            case_id=case_id,
            languages=languages,
        )

    def ingest_sync(
        self,
        *,
        path: Path,
        company_id: str,
        case_id: str = "",
        languages: list[str] | None = None,
    ) -> DocumentManifest:
        resolved = path.expanduser().resolve()
        suffix = validate_document(resolved)
        digest = _sha256(resolved)
        company = company_id.strip() or "default"
        document_id = hashlib.sha256(f"{company}:{digest}".encode()).hexdigest()[:32]
        parsed = parse_document(
            resolved,
            document_id,
            languages or ["vie", "eng"],
            excel_rows_per_unit=self.excel_rows_per_unit,
        )
        return self._persist(
            resolved,
            suffix,
            digest,
            document_id,
            company,
            case_id or document_id,
            parsed,
        )

    def register_extraction(
        self,
        *,
        path: Path,
        company_id: str,
        case_id: str,
        source: str,
        units: Iterable[dict[str, Any]],
        confidence: float,
    ) -> DocumentManifest:
        """Persist units already produced by the compatibility extractor."""
        resolved = path.expanduser().resolve()
        suffix = validate_document(resolved)
        digest = _sha256(resolved)
        company = company_id.strip() or "default"
        document_id = hashlib.sha256(f"{company}:{digest}".encode()).hexdigest()[:32]
        parsed_units = []
        for ordinal, payload in enumerate(units, 1):
            item = dict(payload)
            kind = str(item.get("kind") or "unit")
            locator = str(item.get("locator") or f"unit={ordinal}")
            parsed_units.append(
                _new_unit(
                    document_id,
                    kind,
                    int(item.get("ordinal") or ordinal),
                    locator,
                    text=str(item.get("text") or ""),
                    structured=dict(item.get("structured") or {}),
                    status=str(item.get("status") or "processed"),
                    error=str(item.get("error") or ""),
                    warnings=list(item.get("warnings") or []),
                )
            )
        return self._persist(
            resolved,
            suffix,
            digest,
            document_id,
            company,
            case_id or document_id,
            ParsedDocument(
                source,
                parsed_units,
                "",
                confidence,
                {"extraction_confidence": confidence},
            ),
        )

    def _persist(
        self,
        resolved: Path,
        suffix: str,
        digest: str,
        document_id: str,
        company: str,
        case_id: str,
        parsed: ParsedDocument,
    ) -> DocumentManifest:
        document_dir = self.root / document_id
        manifest_path = document_dir / "manifest.json"
        if manifest_path.is_file():
            manifest = self._read_manifest(manifest_path)
            case_ids = list(manifest.metadata.get("case_ids") or [manifest.case_id])
            if case_id not in case_ids:
                case_ids.append(case_id)
                manifest.metadata["case_ids"] = case_ids
                self._write_manifest(manifest_path, manifest)
            return manifest

        document_dir.mkdir(parents=True, exist_ok=True)
        units_path = document_dir / "units.jsonl"
        failures: list[str] = []
        warnings: list[str] = []
        processed = 0
        total = 0
        temp_units = units_path.with_suffix(".tmp")
        try:
            with temp_units.open("w", encoding="utf-8", newline="\n") as stream:
                for unit in parsed.units:
                    total += 1
                    stream.write(
                        json.dumps(
                            asdict(unit),
                            ensure_ascii=False,
                            default=str,
                        )
                        + "\n"
                    )
                    warnings.extend(unit.warnings)
                    if unit.status == "processed":
                        processed += 1
                    else:
                        failures.append(unit.locator)
                    if self.progress_callback:
                        self.progress_callback(unit, total)
            os.replace(temp_units, units_path)
        finally:
            temp_units.unlink(missing_ok=True)
        company_key = hashlib.sha256(company.encode()).hexdigest()[:16]
        raw_key = f"{company_key}/{document_id}/raw{suffix}"
        units_key = f"{company_key}/{document_id}/units.jsonl"
        self.object_store.put_file(raw_key, resolved, sha256=digest)
        self.object_store.put_file(units_key, units_path)
        metadata = {
            **parsed.metadata,
            "extraction_schema_version": DOCUMENT_EXTRACTION_SCHEMA_VERSION,
            "excel_rows_per_unit": self.excel_rows_per_unit,
            "units_object_key": units_key,
            "extraction_confidence": parsed.confidence,
            "case_ids": [case_id],
        }
        if isinstance(parsed.compatibility_content, dict):
            analysis_path = document_dir / "analysis.json"
            analysis_temp = analysis_path.with_suffix(".tmp")
            try:
                with analysis_temp.open(
                    "w", encoding="utf-8", newline="\n"
                ) as stream:
                    json.dump(
                        parsed.compatibility_content,
                        stream,
                        ensure_ascii=False,
                        default=str,
                    )
                    stream.write("\n")
                os.replace(analysis_temp, analysis_path)
            finally:
                analysis_temp.unlink(missing_ok=True)
            metadata["analysis_uri"] = str(analysis_path)
        coverage = CoverageReport(
            total=total,
            processed=processed,
            failed=len(failures),
            failed_units=tuple(failures),
            warnings=tuple(dict.fromkeys(warnings)),
        )
        manifest = DocumentManifest(
            id=document_id,
            company_id=company,
            case_id=case_id,
            name=resolved.name,
            suffix=suffix,
            sha256=digest,
            size_bytes=resolved.stat().st_size,
            source=parsed.source,
            raw_uri=f"object:{raw_key}",
            units_uri=str(units_path),
            created_at=datetime.now(timezone.utc).isoformat(),
            coverage=coverage,
            metadata=metadata,
        )
        self._write_manifest(manifest_path, manifest)
        return manifest

    async def ingest_case(
        self,
        *,
        paths: Iterable[Path],
        company_id: str,
        case_id: str = "",
        languages: list[str] | None = None,
    ) -> CaseManifest:
        path_list = [path.expanduser().resolve() for path in paths]
        size = sum(path.stat().st_size for path in path_list)
        if size > MAX_DOCUMENT_BYTES:
            raise DocumentValidationError("Case exceeds the 500 MB limit")
        resolved_case = case_id or hashlib.sha256(
            (company_id + ":" + ":".join(_sha256(path) for path in path_list)).encode()
        ).hexdigest()[:32]
        manifests = [
            await self.ingest(
                path=path,
                company_id=company_id,
                case_id=resolved_case,
                languages=languages,
            )
            for path in path_list
        ]
        return CaseManifest(
            resolved_case,
            company_id or "default",
            tuple(item.id for item in manifests),
            size,
            all(item.coverage.complete for item in manifests),
        )

    async def query(
        self,
        *,
        document_id: str,
        company_id: str,
        query: str,
        limit: int = 20,
    ) -> list[DocumentUnit]:
        return await asyncio.to_thread(
            self.query_sync,
            document_id=document_id,
            company_id=company_id,
            query=query,
            limit=limit,
        )

    def query_sync(
        self,
        *,
        document_id: str,
        company_id: str,
        query: str,
        limit: int = 20,
    ) -> list[DocumentUnit]:
        terms = {token.casefold() for token in _TOKEN.findall(query)}
        scored: list[tuple[int, int, DocumentUnit]] = []
        for unit in self.iter_units(document_id, company_id=company_id):
            corpus = f"{unit.locator} {unit.text} {json.dumps(unit.structured, ensure_ascii=False)}"
            folded = corpus.casefold()
            score = sum(folded.count(term) for term in terms)
            if score or not terms:
                scored.append((score, -unit.ordinal, unit))
        scored.sort(key=lambda item: (item[0], item[1]), reverse=True)
        return [item[2] for item in scored[: max(1, limit)]]

    async def coverage(self, document_id: str, *, company_id: str) -> CoverageReport:
        return (
            await asyncio.to_thread(
                self.get_manifest,
                document_id,
                company_id=company_id,
            )
        ).coverage

    def get_manifest(self, document_id: str, *, company_id: str) -> DocumentManifest:
        path = self.root / document_id / "manifest.json"
        if not path.is_file():
            raise KeyError(document_id)
        manifest = self._read_manifest(path)
        if manifest.company_id != company_id:
            raise KeyError(document_id)
        return manifest

    def restore_manifest(
        self,
        manifest: DocumentManifest | Mapping[str, Any],
        units_path: Path,
        *,
        analysis_path: Path | None = None,
    ) -> DocumentManifest:
        """Materialize a distributed manifest and unit stream into the local index."""
        restored = (
            replace(manifest, metadata=dict(manifest.metadata))
            if isinstance(manifest, DocumentManifest)
            else self._manifest_from_payload(dict(manifest))
        )
        if not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9._-]{0,127}", restored.id):
            raise DocumentValidationError("Unsafe document id in restored manifest")
        if not re.fullmatch(r"[0-9a-fA-F]{64}", restored.sha256):
            raise DocumentValidationError("Invalid SHA-256 in restored manifest")
        if not restored.company_id.strip() or not restored.case_id.strip():
            raise DocumentValidationError(
                "Restored manifest requires company_id and case_id"
            )
        source = units_path.expanduser().resolve(strict=True)
        if not source.is_file():
            raise DocumentValidationError("Restored units source is not a file")
        document_dir = self.root / restored.id
        manifest_path = document_dir / "manifest.json"
        if manifest_path.is_file():
            existing = self._read_manifest(manifest_path)
            if (
                existing.company_id != restored.company_id
                or existing.sha256 != restored.sha256
            ):
                raise DocumentValidationError(
                    "Restored document id conflicts with the local index"
                )
        document_dir.mkdir(parents=True, exist_ok=True)
        units_target = document_dir / "units.jsonl"
        descriptor, temporary_name = tempfile.mkstemp(
            prefix=".units-", suffix=".tmp", dir=document_dir
        )
        total = processed = 0
        try:
            with (
                source.open(encoding="utf-8") as input_stream,
                os.fdopen(descriptor, "w", encoding="utf-8", newline="\n") as output,
            ):
                for line_number, line in enumerate(input_stream, 1):
                    try:
                        payload = json.loads(line)
                        unit = DocumentUnit(**payload)
                    except (json.JSONDecodeError, TypeError) as exc:
                        raise DocumentValidationError(
                            f"Invalid restored document unit at line {line_number}"
                        ) from exc
                    if unit.document_id != restored.id:
                        raise DocumentValidationError(
                            f"Restored document unit mismatch at line {line_number}"
                        )
                    if (
                        not unit.id
                        or not unit.kind
                        or not unit.locator
                        or unit.ordinal < 0
                    ):
                        raise DocumentValidationError(
                            f"Invalid restored document unit at line {line_number}"
                        )
                    total += 1
                    processed += int(unit.status == "processed")
                    output.write(line)
                    if not line.endswith("\n"):
                        output.write("\n")
                output.flush()
                os.fsync(output.fileno())
            unresolved = total - processed
            if (
                total != restored.coverage.total
                or processed != restored.coverage.processed
                or unresolved
                != restored.coverage.failed + restored.coverage.skipped
            ):
                raise DocumentValidationError(
                    "Restored unit stream does not match manifest coverage"
                )
            os.replace(temporary_name, units_target)
        finally:
            Path(temporary_name).unlink(missing_ok=True)
        metadata = dict(restored.metadata)
        if analysis_path is None:
            metadata.pop("analysis_uri", None)
        else:
            analysis_source = analysis_path.expanduser().resolve(strict=True)
            try:
                analysis = json.loads(analysis_source.read_text(encoding="utf-8"))
            except json.JSONDecodeError as exc:
                raise DocumentValidationError(
                    "Restored analysis payload is not valid JSON"
                ) from exc
            if not isinstance(analysis, dict):
                raise DocumentValidationError(
                    "Restored analysis payload must be an object"
                )
            analysis_target = document_dir / "analysis.json"
            analysis_temp = analysis_target.with_suffix(".tmp")
            analysis_temp.write_text(
                json.dumps(analysis, ensure_ascii=False, default=str) + "\n",
                encoding="utf-8",
            )
            os.replace(analysis_temp, analysis_target)
            metadata["analysis_uri"] = str(analysis_target)
        restored = replace(
            restored,
            units_uri=str(units_target),
            metadata=metadata,
        )
        self._write_manifest(manifest_path, restored)
        return restored

    def iter_units(
        self,
        document_id: str,
        *,
        company_id: str,
    ) -> Iterator[DocumentUnit]:
        self.get_manifest(document_id, company_id=company_id)
        path = self.root / document_id / "units.jsonl"
        if not path.is_file():
            raise KeyError(document_id)
        with path.open(encoding="utf-8") as stream:
            for line in stream:
                payload = json.loads(line)
                yield DocumentUnit(**payload)

    def analysis_content(
        self,
        document_id: str,
        *,
        company_id: str,
    ) -> dict[str, Any] | str:
        manifest = self.get_manifest(document_id, company_id=company_id)
        path = Path(str(manifest.metadata.get("analysis_uri") or ""))
        if path.is_file():
            return json.loads(path.read_text(encoding="utf-8"))
        return {
            "metadata": manifest.metadata,
            "data": {
                "canonical_metrics": {},
                "coverage": {
                    "total": manifest.coverage.total,
                    "processed": manifest.coverage.processed,
                    "failed": manifest.coverage.failed,
                    "complete": manifest.coverage.complete,
                },
            },
        }

    def prompt_partitions(
        self,
        document_ids: Iterable[str],
        *,
        company_id: str,
        max_chars: int = PROMPT_PARTITION_CHARS,
    ) -> Iterator[str]:
        if max_chars < 1_000:
            raise ValueError("max_chars must be at least 1000")
        batch: list[str] = []
        size = 0
        for document_id in document_ids:
            manifest = self.get_manifest(document_id, company_id=company_id)
            for unit in self.iter_units(document_id, company_id=company_id):
                payload = json.dumps(
                    {
                        "file": manifest.name,
                        "source_id": unit.id,
                        "locator": unit.locator,
                        "kind": unit.kind,
                        "status": unit.status,
                        "text": unit.text,
                        "structured": unit.structured,
                        "warnings": unit.warnings,
                        "error": unit.error,
                    },
                    ensure_ascii=False,
                    default=str,
                )
                for part in _split_complete(payload, max_chars):
                    if batch and size + len(part) + 1 > max_chars:
                        yield "\n".join(batch)
                        batch, size = [], 0
                    batch.append(part)
                    size += len(part) + 1
        if batch:
            yield "\n".join(batch)

    @staticmethod
    def _write_manifest(path: Path, manifest: DocumentManifest) -> None:
        payload = asdict(manifest)
        temp = path.with_suffix(".tmp")
        temp.write_text(
            json.dumps(payload, ensure_ascii=False, indent=2, default=str) + "\n",
            encoding="utf-8",
        )
        os.replace(temp, path)

    @staticmethod
    def _read_manifest(path: Path) -> DocumentManifest:
        payload = json.loads(path.read_text(encoding="utf-8"))
        return DocumentService._manifest_from_payload(payload)

    @staticmethod
    def _manifest_from_payload(payload: Mapping[str, Any]) -> DocumentManifest:
        payload = dict(payload)
        payload["coverage"] = CoverageReport(
            **{
                **payload["coverage"],
                "failed_units": tuple(payload["coverage"].get("failed_units", [])),
                "warnings": tuple(payload["coverage"].get("warnings", [])),
            }
        )
        return DocumentManifest(**payload)


def document_service_from_settings(settings: Mapping[str, Any]):
    data_plane = settings.get("data_plane")
    configured = data_plane if isinstance(data_plane, Mapping) else {}
    if (
        configured.get("enabled")
        or configured.get("postgres_dsn")
        or os.getenv("TAXSENTRY_POSTGRES_DSN")
    ):
        from .data_plane import DistributedDocumentService

        return DistributedDocumentService.from_settings(settings)
    documents = settings.get("documents")
    configured_documents = documents if isinstance(documents, Mapping) else {}
    return DocumentService(
        excel_rows_per_unit=int(
            configured_documents.get("excel_rows_per_unit") or EXCEL_ROWS_PER_UNIT
        )
    )


def _split_complete(text: str, limit: int) -> Iterator[str]:
    """Split without dropping characters; continuation markers preserve lineage."""
    if len(text) <= limit:
        yield text
        return
    body_limit = max(1, limit - 48)
    count = (len(text) + body_limit - 1) // body_limit
    for index in range(count):
        start = index * body_limit
        yield f"[fragment {index + 1}/{count}]\n{text[start:start + body_limit]}"


def pack_complete(values: Iterable[str], max_chars: int) -> Iterator[str]:
    batch: list[str] = []
    size = 0
    for value in values:
        for fragment in _split_complete(value, max_chars):
            if batch and size + len(fragment) + 1 > max_chars:
                yield "\n".join(batch)
                batch, size = [], 0
            batch.append(fragment)
            size += len(fragment) + 1
    if batch:
        yield "\n".join(batch)


def _excel_document(
    path: Path,
    document_id: str,
    *,
    rows_per_unit: int = EXCEL_ROWS_PER_UNIT,
) -> ParsedDocument:
    from openpyxl import load_workbook

    workbook = load_workbook(
        path,
        read_only=True,
        data_only=False,
        keep_links=False,
        keep_vba=path.suffix.casefold() == ".xlsm",
    )
    inventory: list[dict[str, Any]] = []
    defined_names: list[dict[str, str]] = []
    with zipfile.ZipFile(path) as archive:
        chart_sources = _chart_sources(archive)
        external_links = [
            name
            for name in archive.namelist()
            if name.startswith("xl/externalLinks/") and name.endswith(".xml")
        ]
        for value in workbook.defined_names.values():
            defined_names.append(
                {
                    "name": str(getattr(value, "name", "")),
                    "reference": str(getattr(value, "attr_text", "") or ""),
                }
            )
        for worksheet in workbook.worksheets:
            xml_path = str(getattr(worksheet, "_worksheet_path", "")).lstrip("/")
            xml_meta = (
                _sheet_xml_metadata(archive, xml_path) if xml_path in archive.namelist() else {}
            )
            inventory.append(
                {
                    "name": worksheet.title,
                    "state": worksheet.sheet_state,
                    "rows": worksheet.max_row,
                    "columns": worksheet.max_column,
                    **xml_meta,
                }
            )
    workbook.close()
    metadata = {
        "sheets": inventory,
        "defined_names": defined_names,
        "chart_sources": chart_sources,
        "external_links": external_links,
        "macros_present": path.suffix.casefold() == ".xlsm",
        "macros_executed": False,
        "external_links_opened": False,
    }

    def units() -> Iterator[DocumentUnit]:
        yield _new_unit(
            document_id,
            "workbook",
            0,
            "workbook",
            text="; ".join(
                f"{item['name']} ({item['state']}, {item['rows']}x{item['columns']})"
                for item in inventory
            ),
            structured=metadata,
            warnings=(
                ["Workbook contains formulas without cached values"]
                if any(item.get("formulas_without_cached_values") for item in inventory)
                else []
            ),
        )
        streaming = load_workbook(
            path,
            read_only=True,
            data_only=False,
            keep_links=False,
            keep_vba=path.suffix.casefold() == ".xlsm",
        )
        ordinal = 1
        try:
            for worksheet in streaming.worksheets:
                rows: list[dict[str, Any]] = []
                start_row = 1
                sheet_had_rows = False
                for row_number, row in enumerate(
                    worksheet.iter_rows(values_only=False), 1
                ):
                    cells: list[dict[str, Any]] = []
                    dependencies: list[dict[str, Any]] = []
                    for cell in row:
                        value = _json_value(cell.value)
                        if value is None:
                            continue
                        cells.append({"cell": cell.coordinate, "value": value})
                        if isinstance(value, str) and value.startswith("="):
                            dependencies.append(
                                {
                                    "cell": cell.coordinate,
                                    "references": [
                                        {
                                            "sheet": match.group(1)
                                            or (match.group(2) or "").strip()
                                            or worksheet.title,
                                            "cell": (
                                                f"{match.group(3)}{match.group(4)}"
                                            ),
                                        }
                                        for match in _FORMULA_REF.finditer(
                                            value
                                        )
                                    ],
                                }
                            )
                    if cells:
                        rows.append(
                            {
                                "row": row_number,
                                "cells": cells,
                                **(
                                    {"formula_dependencies": dependencies}
                                    if dependencies
                                    else {}
                                ),
                            }
                        )
                        sheet_had_rows = True
                    if len(rows) >= rows_per_unit:
                        locator = (
                            f"sheet={worksheet.title}!rows={start_row}:{row_number}"
                        )
                        yield _new_unit(
                            document_id,
                            "sheet_rows",
                            ordinal,
                            locator,
                            text=_excel_rows_text(rows),
                            structured={"sheet": worksheet.title, "rows": rows},
                        )
                        ordinal += 1
                        rows, start_row = [], row_number + 1
                if rows:
                    locator = (
                        f"sheet={worksheet.title}!rows={start_row}:{rows[-1]['row']}"
                    )
                    yield _new_unit(
                        document_id,
                        "sheet_rows",
                        ordinal,
                        locator,
                        text=_excel_rows_text(rows),
                        structured={"sheet": worksheet.title, "rows": rows},
                    )
                    ordinal += 1
                if not sheet_had_rows:
                    yield _new_unit(
                        document_id,
                        "sheet_rows",
                        ordinal,
                        f"sheet={worksheet.title}!rows=empty",
                        structured={"sheet": worksheet.title, "rows": []},
                    )
                    ordinal += 1
        finally:
            streaming.close()

    compatibility: dict[str, Any]
    total_cells = sum(
        int(item["rows"]) * int(item["columns"]) for item in inventory
    )
    if path.stat().st_size <= 25 * 1024 * 1024 and total_cells <= 250_000:
        from .core.excel_parser import TaxSentryParser

        parser = TaxSentryParser(str(path))
        parser.load()
        compatibility = (
            json.loads(parser.export_json())
            if parser.has_meaningful_data()
            else {"metadata": metadata, "data": {}}
        )
    else:
        compatibility = {
            "metadata": metadata,
            "data": {
                "canonical_metrics": {},
                "sheets": inventory,
                "coverage": "Use DocumentService units for complete row coverage",
            },
        }
    return ParsedDocument(
        path.suffix.casefold().lstrip("."),
        units(),
        compatibility,
        1.0,
        metadata,
    )


def _excel_rows_text(rows: list[dict[str, Any]]) -> str:
    lines = []
    for row in rows:
        cells = " | ".join(
            f"{cell['cell']}={_text_value(cell['value'])}" for cell in row["cells"]
        )
        lines.append(f"row {row['row']}: {cells}")
    return "\n".join(lines)


def _sheet_xml_metadata(archive: zipfile.ZipFile, member: str) -> dict[str, Any]:
    dimension = ""
    merged: list[str] = []
    formula_count = 0
    missing_cached = 0
    with archive.open(member) as stream:
        for _, element in ElementTree.iterparse(stream, events=("end",)):
            local = element.tag.rsplit("}", 1)[-1]
            if local == "dimension":
                dimension = element.attrib.get("ref", "")
            elif local == "mergeCell":
                merged.append(element.attrib.get("ref", ""))
            elif local == "c":
                formula = next(
                    (
                        child
                        for child in element
                        if child.tag.rsplit("}", 1)[-1] == "f"
                    ),
                    None,
                )
                if formula is not None:
                    formula_count += 1
                    cached = next(
                        (
                            child
                            for child in element
                            if child.tag.rsplit("}", 1)[-1] == "v"
                        ),
                        None,
                    )
                    if cached is None or cached.text in {None, ""}:
                        missing_cached += 1
                element.clear()
    return {
        "used_range": dimension,
        "merged_cells": merged,
        "formula_count": formula_count,
        "formulas_without_cached_values": missing_cached,
    }


def _chart_sources(archive: zipfile.ZipFile) -> list[str]:
    sources: list[str] = []
    for name in archive.namelist():
        if not (name.startswith("xl/charts/chart") and name.endswith(".xml")):
            continue
        root = ElementTree.fromstring(archive.read(name))
        for element in root.iter():
            if element.tag.rsplit("}", 1)[-1] == "f" and element.text:
                sources.append(element.text)
    return list(dict.fromkeys(sources))


def _word_document(
    path: Path, document_id: str, languages: list[str]
) -> ParsedDocument:
    units: list[DocumentUnit] = []
    ordinal = 1
    with zipfile.ZipFile(path) as archive:
        root = ElementTree.fromstring(archive.read("word/document.xml"))
        body = next(
            (element for element in root.iter() if element.tag.rsplit("}", 1)[-1] == "body"),
            root,
        )
        section = ""
        for child in body:
            local = child.tag.rsplit("}", 1)[-1]
            if local == "p":
                text = _xml_text(child)
                style = _paragraph_style(child)
                if style.casefold().startswith("heading"):
                    section = text
                if text:
                    units.append(
                        _new_unit(
                            document_id,
                            "paragraph",
                            ordinal,
                            f"section={section or 'body'};paragraph={ordinal}",
                            text=text,
                            structured={
                                "style": style,
                                "section": section,
                                "tracked_change": _has_tracked_change(child),
                            },
                        )
                    )
                    ordinal += 1
            elif local == "tbl":
                rows = _xml_table(child)
                units.append(
                    _new_unit(
                        document_id,
                        "table",
                        ordinal,
                        f"section={section or 'body'};table={ordinal}",
                        text="\n".join(" | ".join(row) for row in rows),
                        structured={"section": section, "rows": rows},
                    )
                )
                ordinal += 1
        if not units:
            text = _xml_text(root)
            if text:
                units.append(
                    _new_unit(
                        document_id,
                        "paragraph",
                        ordinal,
                        "section=body;paragraph=1",
                        text=text,
                    )
                )
                ordinal += 1
        for prefix, kind in (
            ("word/header", "header"),
            ("word/footer", "footer"),
            ("word/footnotes", "footnote"),
            ("word/endnotes", "endnote"),
            ("word/comments", "comment"),
        ):
            for name in sorted(
                item
                for item in archive.namelist()
                if item.startswith(prefix) and item.endswith(".xml")
            ):
                xml = ElementTree.fromstring(archive.read(name))
                for element in xml.iter():
                    if element.tag.rsplit("}", 1)[-1] not in {"p", kind}:
                        continue
                    text = _xml_text(element)
                    if text:
                        units.append(
                            _new_unit(
                                document_id,
                                kind,
                                ordinal,
                                f"{kind}={Path(name).stem};item={ordinal}",
                                text=text,
                                structured={"part": name},
                            )
                        )
                        ordinal += 1
        for name in sorted(
            item for item in archive.namelist() if item.startswith("word/media/")
        ):
            try:
                text, confidence, boxes = _ocr_bytes(archive.read(name), languages)
                units.append(
                    _new_unit(
                        document_id,
                        "image",
                        ordinal,
                        f"image={Path(name).name}",
                        text=text,
                        structured={"confidence": confidence, "words": boxes},
                    )
                )
            except Exception as exc:
                units.append(
                    _new_unit(
                        document_id,
                        "image",
                        ordinal,
                        f"image={Path(name).name}",
                        status="failed",
                        error=str(exc),
                    )
                )
            ordinal += 1
    text = "\n\n".join(unit.text for unit in units if unit.text)
    confidence = _average_confidence(units)
    return ParsedDocument("docx", units, text, confidence)


def _paragraph_style(element: ElementTree.Element) -> str:
    for child in element.iter():
        if child.tag.rsplit("}", 1)[-1] == "pStyle":
            return next(
                (value for key, value in child.attrib.items() if key.endswith("}val")),
                child.attrib.get("val", ""),
            )
    return ""


def _has_tracked_change(element: ElementTree.Element) -> bool:
    return any(
        child.tag.rsplit("}", 1)[-1] in {"ins", "del", "moveFrom", "moveTo"}
        for child in element.iter()
    )


def _xml_text(element: ElementTree.Element) -> str:
    return " ".join(
        str(child.text).strip()
        for child in element.iter()
        if child.tag.rsplit("}", 1)[-1] in {"t", "delText"}
        and child.text
        and str(child.text).strip()
    )


def _xml_table(element: ElementTree.Element) -> list[list[str]]:
    rows: list[list[str]] = []
    for row in element:
        if row.tag.rsplit("}", 1)[-1] != "tr":
            continue
        rows.append(
            [
                _xml_text(cell)
                for cell in row
                if cell.tag.rsplit("}", 1)[-1] == "tc"
            ]
        )
    return rows


def _powerpoint_document(
    path: Path, document_id: str, languages: list[str]
) -> ParsedDocument:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    try:
        presentation = Presentation(path)
    except Exception:
        return _powerpoint_xml_fallback(path, document_id)
    units: list[DocumentUnit] = []
    confidences: list[float] = []
    for slide_number, slide in enumerate(presentation.slides, 1):
        shape_items: list[dict[str, Any]] = []
        text_parts: list[str] = []
        warnings: list[str] = []
        for shape_index, shape in enumerate(slide.shapes, 1):
            item: dict[str, Any] = {
                "shape": shape_index,
                "name": str(getattr(shape, "name", "")),
                "left": int(getattr(shape, "left", 0)),
                "top": int(getattr(shape, "top", 0)),
                "width": int(getattr(shape, "width", 0)),
                "height": int(getattr(shape, "height", 0)),
            }
            if getattr(shape, "has_text_frame", False):
                item["text"] = shape.text
                text_parts.append(shape.text)
            if getattr(shape, "has_table", False):
                rows = [
                    [cell.text for cell in row.cells] for row in shape.table.rows
                ]
                item["table"] = rows
                text_parts.extend(" | ".join(row) for row in rows)
            if getattr(shape, "has_chart", False):
                chart = []
                for series in shape.chart.series:
                    values = [_json_value(value) for value in series.values]
                    chart.append({"name": str(series.name), "values": values})
                    text_parts.append(f"{series.name}: {values}")
                item["chart_series"] = chart
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                try:
                    text, confidence, boxes = _ocr_bytes(shape.image.blob, languages)
                    item["image_ocr"] = {
                        "text": text,
                        "confidence": confidence,
                        "words": boxes,
                    }
                    text_parts.append(text)
                    confidences.append(confidence)
                except Exception as exc:
                    warnings.append(f"shape {shape_index} OCR failed: {exc}")
            shape_items.append(item)
        notes = ""
        try:
            notes = slide.notes_slide.notes_text_frame.text
        except (AttributeError, ValueError):
            pass
        if notes:
            text_parts.append(notes)
        units.append(
            _new_unit(
                document_id,
                "slide",
                slide_number,
                f"slide={slide_number}",
                text="\n".join(part for part in text_parts if part),
                structured={"shapes": shape_items, "speaker_notes": notes},
                warnings=warnings,
            )
        )
    with zipfile.ZipFile(path) as archive:
        ordinal = len(units) + 1
        for name in sorted(
            item
            for item in archive.namelist()
            if (
                item.startswith("ppt/comments/comment")
                or item.startswith("ppt/slideMasters/slideMaster")
                or item.startswith("ppt/theme/theme")
            )
            and item.endswith(".xml")
        ):
            kind = (
                "comment"
                if "/comments/" in name
                else ("slide_master" if "/slideMasters/" in name else "theme")
            )
            text = _xml_text(ElementTree.fromstring(archive.read(name)))
            units.append(
                _new_unit(
                    document_id,
                    kind,
                    ordinal,
                    f"{kind}={Path(name).stem}",
                    text=text,
                    structured={"part": name},
                )
            )
            ordinal += 1
    text = "\n\n".join(unit.text for unit in units if unit.text)
    return ParsedDocument(
        "pptx",
        units,
        text,
        sum(confidences) / len(confidences) if confidences else (1.0 if text else 0.0),
        {"slide_count": len(presentation.slides)},
    )


def _powerpoint_xml_fallback(path: Path, document_id: str) -> ParsedDocument:
    units: list[DocumentUnit] = []
    with zipfile.ZipFile(path) as archive:
        slides = sorted(
            (
                name
                for name in archive.namelist()
                if name.startswith("ppt/slides/slide") and name.endswith(".xml")
            ),
            key=_natural_key,
        )
        for ordinal, name in enumerate(slides, 1):
            text = _xml_text(ElementTree.fromstring(archive.read(name)))
            units.append(
                _new_unit(
                    document_id,
                    "slide",
                    ordinal,
                    f"slide={ordinal}",
                    text=text,
                    structured={"part": name},
                )
            )
    text = "\n\n".join(unit.text for unit in units if unit.text)
    return ParsedDocument(
        "pptx",
        units,
        text,
        1.0 if text else 0.0,
        {"slide_count": len(units)},
    )


def _natural_key(value: str) -> list[int | str]:
    return [
        int(part) if part.isdigit() else part.casefold()
        for part in re.split(r"([0-9]+)", value)
    ]


def _pdf_document(
    path: Path, document_id: str, languages: list[str]
) -> ParsedDocument:
    import fitz
    import pdfplumber
    from PIL import Image

    units: list[DocumentUnit] = []
    confidences: list[float] = []
    with pdfplumber.open(path) as plumber, fitz.open(path) as fitz_document:
        for page_number, page in enumerate(plumber.pages, 1):
            warnings: list[str] = []
            status, error = "processed", ""
            text = page.extract_text() or ""
            words = [
                {
                    "text": word.get("text", ""),
                    "bbox": [
                        word.get("x0"),
                        word.get("top"),
                        word.get("x1"),
                        word.get("bottom"),
                    ],
                }
                for word in page.extract_words()
            ]
            source = "text"
            confidence = 1.0 if text else 0.0
            if len(text.strip()) < 80:
                try:
                    pixmap = fitz_document[page_number - 1].get_pixmap(
                        matrix=fitz.Matrix(2, 2), alpha=False
                    )
                    image = Image.frombytes(
                        "RGB", (pixmap.width, pixmap.height), pixmap.samples
                    )
                    ocr_text, confidence, ocr_words = _ocr_image(image, languages)
                    if ocr_text.strip():
                        text = "\n".join(part for part in (text, ocr_text) if part)
                        words.extend(ocr_words)
                        source = "hybrid-ocr" if page.extract_text() else "ocr"
                    else:
                        warnings.append("OCR returned no text")
                except Exception as exc:
                    status, error = "failed", f"OCR failed: {exc}"
            confidences.append(confidence)
            try:
                tables = page.extract_tables() or []
            except Exception as exc:
                tables = []
                warnings.append(f"table extraction failed: {exc}")
            units.append(
                _new_unit(
                    document_id,
                    "page",
                    page_number,
                    f"page={page_number}",
                    text=text,
                    structured={
                        "text_source": source,
                        "confidence": confidence,
                        "width": page.width,
                        "height": page.height,
                        "words": words,
                        "tables": tables,
                        "images": page.images,
                    },
                    status=status,
                    error=error,
                    warnings=warnings,
                )
            )
    text = "\n\n".join(unit.text for unit in units if unit.text)
    return ParsedDocument(
        "pdf-pages",
        units,
        text,
        sum(confidences) / len(confidences) if confidences else 0.0,
        {"page_count": len(units)},
    )


def _image_document(
    path: Path, document_id: str, languages: list[str]
) -> ParsedDocument:
    from PIL import Image

    text, confidence, words = _ocr_image(Image.open(path), languages)
    unit = _new_unit(
        document_id,
        "image",
        1,
        "image=1",
        text=text,
        structured={"confidence": confidence, "words": words},
    )
    return ParsedDocument("ocr", [unit], text, confidence)


def _ocr_bytes(
    data: bytes, languages: list[str]
) -> tuple[str, float, list[dict[str, Any]]]:
    from io import BytesIO

    from PIL import Image

    return _ocr_image(Image.open(BytesIO(data)), languages)


def _ocr_image(
    image: Any, languages: list[str]
) -> tuple[str, float, list[dict[str, Any]]]:
    import pytesseract

    data = pytesseract.image_to_data(
        image.convert("RGB"),
        lang="+".join(languages),
        output_type=pytesseract.Output.DICT,
    )
    words: list[dict[str, Any]] = []
    scores: list[float] = []
    for index, raw in enumerate(data["text"]):
        text = str(raw).strip()
        try:
            score = float(data["conf"][index])
        except (TypeError, ValueError):
            score = -1
        if text:
            words.append(
                {
                    "text": text,
                    "bbox": [
                        int(data["left"][index]),
                        int(data["top"][index]),
                        int(data["left"][index]) + int(data["width"][index]),
                        int(data["top"][index]) + int(data["height"][index]),
                    ],
                }
            )
        if score >= 0:
            scores.append(score)
    return (
        " ".join(item["text"] for item in words),
        sum(scores) / len(scores) / 100 if scores else 0.0,
        words,
    )


def _average_confidence(units: list[DocumentUnit]) -> float:
    scores = [
        float(unit.structured["confidence"])
        for unit in units
        if isinstance(unit.structured.get("confidence"), (int, float))
    ]
    if scores:
        return sum(scores) / len(scores)
    return 1.0 if any(unit.text for unit in units) else 0.0


def _legacy_document(
    path: Path,
    document_id: str,
    languages: list[str],
    *,
    excel_rows_per_unit: int = EXCEL_ROWS_PER_UNIT,
) -> ParsedDocument:
    if (
        os.getenv("TAXSENTRY_OFFICE_SANDBOX") != "1"
        or os.getenv("TAXSENTRY_OFFICE_NETWORK_ISOLATED") != "1"
    ):
        raise DocumentNeedsReview(
            "Network-isolated LibreOffice worker is required for legacy conversion"
        )
    command = shutil.which("soffice") or shutil.which("libreoffice")
    if not command:
        raise DocumentNeedsReview(
            "LibreOffice worker is required to read legacy .doc/.xls/.ppt files"
        )
    target_suffix = {".doc": "docx", ".xls": "xlsx", ".ppt": "pptx"}[
        path.suffix.casefold()
    ]
    with tempfile.TemporaryDirectory(prefix="taxsentry-office-") as folder:
        output = Path(folder)
        profile = output / "profile"
        try:
            result = subprocess.run(
                [
                    command,
                    f"-env:UserInstallation={profile.as_uri()}",
                    "--headless",
                    "--nologo",
                    "--nodefault",
                    "--nofirststartwizard",
                    "--norestore",
                    "--convert-to",
                    target_suffix,
                    "--outdir",
                    str(output),
                    str(path),
                ],
                capture_output=True,
                text=True,
                timeout=90,
                check=False,
                env={
                    **os.environ,
                    "SAL_DISABLE_MACROS": "1",
                },
            )
        except (OSError, subprocess.TimeoutExpired) as exc:
            raise DocumentNeedsReview(f"LibreOffice conversion failed: {exc}") from exc
        converted = output / f"{path.stem}.{target_suffix}"
        if result.returncode or not converted.is_file() or not converted.stat().st_size:
            detail = (result.stderr or result.stdout or "conversion produced no output").strip()
            raise DocumentNeedsReview(
                f"Legacy Office file may be encrypted or invalid: {detail[:300]}"
            )
        parsed = parse_document(
            converted,
            document_id,
            languages,
            excel_rows_per_unit=excel_rows_per_unit,
        )
        parsed.source = f"{path.suffix.casefold()[1:]}->{parsed.source}"
        parsed.metadata.update(
            {
                "converted_in_isolated_profile": True,
                "macros_executed": False,
                "network_policy": "isolated-container-network",
            }
        )
        return parsed
